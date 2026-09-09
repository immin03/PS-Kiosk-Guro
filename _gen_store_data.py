# Generate storeData.js from Excel + rack-sections.json
# 주의: 랙 카테고리 정본은 PS-OS catalog/racks.json 입니다.
#      여기서 만드는 라벨은 화면에 쓰지 않습니다 (src/rackLayout.js 가 정본 사본).
import json, re, sys
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

root = Path(__file__).resolve().parent
sections = json.loads((root / "rack-sections.json").read_text(encoding="utf-8"))
xlsx = Path(r"C:\Users\user\OneDrive\Desktop\구로점 랙별 진열 SKU_260801.xlsx")
if not xlsx.exists():
    # fallback: copy beside repo
    xlsx = root / "구로점 랙별 진열 SKU_260801.xlsx"

# flatten labels
LABELS = {}
for group in ("A", "B", "C", "D"):
    LABELS.update(sections[group])
ALIASES = sections.get("aliases", {})

EMOJIS = {
    "콘드로이친": "🦴", "뼈건강": "🦴", "관절·근육 건강": "💪",
    "오메가3": "👁️", "기억력·인지력": "🧠", "눈건강": "👁️",
    "숙취해소": "🩺", "간건강": "🩺", "수면·스트레스": "😴",
    "이너뷰티": "✨", "여성건강": "🌸", "남성건강": "💪",
    "면역": "🛡️", "항산화": "🛡️", "홍삼": "🌿",
    "종합비타민": "💊", "비타민D": "💊", "비타민C": "💊",
    "단백질·아미노산": "⚡", "활력·에너지": "⚡", "어린이 건강": "👶🏻",
    "혈당·체중관리": "⚖️", "쉐이크·식사대용": "🥤", "디톡스": "🍃",
    "식품": "🍯", "식이섬유·효소": "🦠", "유산균": "🦠",
    "구강건강": "🦷", "구강청결": "🦷", "구강용품": "🪥",
    "모래·패드": "🐾", "펫 건강": "🐾", "펫 미용": "🛁",
    "펫 장난감": "🎾", "펫 위생용품": "🧼", "펫 사료·간식": "🦴",
    "마스크팩": "💧", "토너·세럼·미스트": "💧", "앰플·에센스": "💧",
    "크림": "💧", "아이케어": "👁️", "립 메이크업": "💄",
    "베이스 메이크업": "💄", "선케어": "☀️", "바디로션": "🧴",
    "프레시케어": "🧴", "뷰티 디바이스": "🔌", "샴푸": "💇",
    "트리트먼트·팩": "💇", "헤어케어·염색": "💇", "핸드크림·립밤": "🧴",
    "생리용품": "🧼", "위생용품": "🧼", "헤어 스타일링": "💇",
    "성인용품": "💋", "여성 위생용품": "🧼", "제모용품·스크럽": "✨",
    "탈취·방향제": "🏡", "글로벌 뷰티": "🌍", "클렌징": "🫧",
    "핸드워시": "🧴", "바디케어": "🧴", "제이숲": "💇",
    "티에스": "💇", "꼬꼬도르": "🏡", "브이티 코스메틱": "💧",
    "어나더페이스": "💧", "믹순": "💧", "데싱디바": "💅",
    "풀리오": "🔌", "스파알": "🔌", "에이밍": "🪒", "맨즈케어": "🧔",
    "농협 · 한삼인": "🌿", "대상 웰라이프": "🥤", "네이처가든": "🌿",
    "익스트림": "⚡", "있나요 · 뉴핏": "⚖️", "어린이 간식": "🍪",
    "음료판매대": "🧃", "이벤트존": "🎁",
}


def zone_for(rack: str) -> str:
    if rack in ALIASES:
        rack = ALIASES[rack]
    if not rack:
        return "A"
    letter = rack[0]
    if letter == "A":
        try:
            n = int(re.sub(r"\D", "", rack) or "0")
            if n >= 31:
                return "E"
        except ValueError:
            pass
        return "A"
    if letter in "BCDE":
        return letter
    return "B"


products = []
rack_counts = Counter()
brands = Counter()

wb = load_workbook(xlsx, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
for r in rows[1:]:
    if not r or len(r) < 3:
        continue
    raw_rack = str(r[0] or "").strip()
    sku = str(r[2] or "").strip()
    if not sku:
        continue
    rack = ALIASES.get(raw_rack, raw_rack)
    m = re.match(r"^\[([^\]]+)\]\s*(.*)$", sku)
    if m:
        brand, name = m.group(1).strip(), m.group(2).strip()
    else:
        brand, name = "", sku
    if not name:
        continue
    label = LABELS.get(rack, raw_rack)
    zone = zone_for(rack)
    products.append({
        "name": name,
        "brand": brand,
        "cat": label,
        "rack": rack,
        "zone": zone,
        "en": "",
        "benefit": f"{label} · 섹션 {rack}",
        "spec": "",
        "sale": 0,
    })
    rack_counts[rack] += 1
    if brand:
        brands[brand] += 1

health, beauty, pet, brand_top = [], [], [], []

for rack, label in sections["A"].items():
    n = int(re.sub(r"\D", "", rack))
    entry = {
        "id": rack.lower(),
        "name": label,
        "emoji": EMOJIS.get(label, "📦"),
        "count": rack_counts.get(rack, 0),
        "zone": "E" if n >= 31 else "A",
        "racks": rack,
    }
    (pet if n >= 31 else health).append(entry)

for rack, label in sections["B"].items():
    brand_top.append({
        "id": rack.lower(),
        "name": label,
        "emoji": EMOJIS.get(label, "🏷️"),
        "count": rack_counts.get(rack, 0),
        "zone": "B",
        "racks": rack,
    })

for rack, label in sections["C"].items():
    beauty.append({
        "id": rack.lower(),
        "name": label,
        "emoji": EMOJIS.get(label, "🧴"),
        "count": rack_counts.get(rack, 0),
        "zone": "C",
        "racks": rack,
    })

for rack, label in sections["D"].items():
    beauty.append({
        "id": rack.lower(),
        "name": label,
        "emoji": EMOJIS.get(label, "✨"),
        "count": rack_counts.get(rack, 0),
        "zone": "D",
        "racks": rack,
    })

HEALTH_CATS = health + brand_top
BEAUTY_CATS = beauty
PET_CATS = pet

health_n = sum(1 for p in products if p["zone"] in ("A", "B"))
beauty_n = sum(1 for p in products if p["zone"] in ("C", "D"))
pet_n = sum(1 for p in products if p["zone"] == "E")

TOP_CATS = [
    {"label": "건강기능식품", "sub": f"A·B 섹션 · {health_n}종", "emoji": "💊", "type": "health"},
    {"label": "뷰티 · 라이프 스타일", "sub": f"C·D 섹션 · {beauty_n}종", "emoji": "🎀", "type": "beauty"},
    {"label": "펫", "sub": f"A31–A36 · {pet_n}종", "emoji": "🐾", "type": "pet"},
    {"label": "라이프 스타일", "sub": "프레시케어 · 위생 · 탈취 등", "emoji": "🧴", "type": "life"},
]

ZONES_MAP = [
    {"id": "A", "label": "건강기능식품", "color": "#00978F", "desc": "A1~A30", "catType": "health"},
    {"id": "B", "label": "브랜드 · 간식 · 음료", "color": "#9B7BE0", "desc": "B1~B10", "catType": "health"},
    {"id": "C", "label": "뷰티 · 라이프 스타일", "color": "#EF7BA4", "desc": "C1~C25", "catType": "beauty"},
    {"id": "D", "label": "브랜드 존", "color": "#5BB8E8", "desc": "D1~D15", "catType": "beauty"},
    {"id": "E", "label": "펫", "color": "#FF9F43", "desc": "A31~A36", "catType": "pet"},
]

ALL_BRANDS = sorted(brands.keys(), key=lambda b: (-brands[b], b))

out = root / "src" / "storeData.js"

def dumps(obj):
    return json.dumps(obj, ensure_ascii=False, indent=2)

js = f"""/* Auto-generated from 구로점 랙별 진열 SKU_260801.xlsx + latest floor map.
 * Do not edit by hand — regenerate via _gen_store_data.py
 */

export const HEALTH_CATS = {dumps(HEALTH_CATS)};

export const BEAUTY_CATS = {dumps(BEAUTY_CATS)};

export const PET_CATS = {dumps(PET_CATS)};

export const TOP_CATS = {dumps(TOP_CATS)};

export const ALL_BRANDS = {dumps(ALL_BRANDS)};

export const ZONES_MAP = {dumps(ZONES_MAP)};

export const ALL_PRODUCTS = {dumps(products)};
"""
out.write_text(js, encoding="utf-8")
print("wrote", out, "products", len(products), "brands", len(ALL_BRANDS))
print("health cats", len(HEALTH_CATS), "beauty", len(BEAUTY_CATS), "pet", len(PET_CATS))
print("A10", rack_counts.get("A10", 0), "A12", rack_counts.get("A12", 0), "C2", rack_counts.get("C2", 0), "C19", rack_counts.get("C19", 0))
